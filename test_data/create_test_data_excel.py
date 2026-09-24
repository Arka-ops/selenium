"""
Run this script once to generate test_data.xlsx
Usage: python create_test_data_excel.py
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import os

def create_excel():
    wb = openpyxl.Workbook()

    # ── Sheet 1: Users ──────────────────────────────────────────────────────────
    ws_users = wb.active
    ws_users.title = "Users"
    header_font  = Font(bold=True, color="FFFFFF")
    header_fill  = PatternFill("solid", fgColor="2E86AB")
    center_align = Alignment(horizontal="center")

    user_headers = ["Test ID", "Email", "Password", "Role", "Active"]
    for col, h in enumerate(user_headers, 1):
        cell = ws_users.cell(row=1, column=col, value=h)
        cell.font, cell.fill, cell.alignment = header_font, header_fill, center_align

    user_data = [
        ["TC001", "test@example.com",   "Test@1234",  "registered_user",  "Yes"],
        ["TC002", "admin@example.com",  "Admin@1234", "admin",            "No"],
    ]
    for row_idx, row in enumerate(user_data, 2):
        for col_idx, val in enumerate(row, 1):
            ws_users.cell(row=row_idx, column=col_idx, value=val)

    for col in ws_users.columns:
        ws_users.column_dimensions[col[0].column_letter].width = 22

    # ── Sheet 2: Products ───────────────────────────────────────────────────────
    ws_products = wb.create_sheet("Products")
    prod_fill    = PatternFill("solid", fgColor="A23B72")

    prod_headers = ["Product ID", "Name", "Search Term", "Quantity",
                    "Min Price", "Max Price", "Expected In Cart"]
    for col, h in enumerate(prod_headers, 1):
        cell = ws_products.cell(row=1, column=col, value=h)
        cell.font  = Font(bold=True, color="FFFFFF")
        cell.fill  = prod_fill
        cell.alignment = center_align

    products = [
        ["P001", "iPhone",  "iPhone",  2, 100,  2000, "Yes"],
        ["P002", "MacBook", "MacBook", 1, 500,  5000, "Yes"],
        ["P003", "Samsung", "Samsung", 3, 50,   1500, "Yes"],
    ]
    for row_idx, row in enumerate(products, 2):
        for col_idx, val in enumerate(row, 1):
            ws_products.cell(row=row_idx, column=col_idx, value=val)

    for col in ws_products.columns:
        ws_products.column_dimensions[col[0].column_letter].width = 18

    # ── Sheet 3: Test Cases ─────────────────────────────────────────────────────
    ws_tc = wb.create_sheet("TestCases")
    tc_fill = PatternFill("solid", fgColor="F18F01")

    tc_headers = ["TC ID", "Module", "Test Description", "Priority",
                  "Expected Result", "Status"]
    for col, h in enumerate(tc_headers, 1):
        cell = ws_tc.cell(row=1, column=col, value=h)
        cell.font  = Font(bold=True, color="FFFFFF")
        cell.fill  = tc_fill
        cell.alignment = center_align

    test_cases = [
        ["TC001", "Login",        "Verify login with valid credentials",         "High",   "Login successful",           "Pending"],
        ["TC002", "Search",       "Search for an existing product",               "High",   "Product results displayed",   "Pending"],
        ["TC003", "Add to Cart",  "Add searched product to cart",                 "High",   "Product added to cart",       "Pending"],
        ["TC004", "Update Cart",  "Update product quantity in cart",              "Medium", "Quantity updated",            "Pending"],
        ["TC005", "Verify Cart",  "Verify cart details match expected values",    "High",   "Cart details correct",        "Pending"],
        ["TC006", "Popup/Alert",  "Handle any popup or alert on the page",        "Low",    "Popup handled gracefully",    "Pending"],
        ["TC007", "Screenshot",  "Capture screenshot at each major step",         "Low",    "Screenshots saved",           "Pending"],
        ["TC008", "Report",       "Generate HTML execution report",               "Medium", "Report generated",            "Pending"],
    ]
    for row_idx, row in enumerate(test_cases, 2):
        for col_idx, val in enumerate(row, 1):
            ws_tc.cell(row=row_idx, column=col_idx, value=val)

    for col in ws_tc.columns:
        ws_tc.column_dimensions[col[0].column_letter].width = 30

    out = os.path.join(os.path.dirname(__file__), "test_data.xlsx")
    wb.save(out)
    print(f"✅  test_data.xlsx created → {out}")

if __name__ == "__main__":
    create_excel()
