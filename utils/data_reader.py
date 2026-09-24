"""
utils/data_reader.py
Reads test data from JSON or Excel (openpyxl).
"""

import json
import os
from typing import Any, Dict, List, Optional

from config.config import TEST_DATA_JSON, TEST_DATA_EXCEL


# ── JSON helpers ──────────────────────────────────────────────────────────────

def read_json(filepath: str = TEST_DATA_JSON) -> Dict[str, Any]:
    with open(filepath, "r", encoding="utf-8") as fh:
        return json.load(fh)


def get_user_from_json(user_id: str = "TC001") -> Optional[Dict]:
    data = read_json()
    for user in data.get("users", []):
        if user["id"] == user_id:
            return user
    return None


def get_product_from_json(product_id: str = "P001") -> Optional[Dict]:
    data = read_json()
    for product in data.get("products", []):
        if product["id"] == product_id:
            return product
    return None


# ── Excel helpers ─────────────────────────────────────────────────────────────

def _sheet_to_dicts(wb_path: str, sheet_name: str) -> List[Dict[str, Any]]:
    """
    Read an Excel sheet and return a list of row-dicts keyed by header row.
    Requires openpyxl.
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required for Excel reading: pip install openpyxl")

    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else f"col_{i}"
               for i, h in enumerate(rows[0])]
    return [dict(zip(headers, row)) for row in rows[1:] if any(v is not None for v in row)]


def get_users_from_excel(filepath: str = TEST_DATA_EXCEL) -> List[Dict]:
    return _sheet_to_dicts(filepath, "Users")


def get_products_from_excel(filepath: str = TEST_DATA_EXCEL) -> List[Dict]:
    return _sheet_to_dicts(filepath, "Products")


def get_test_cases_from_excel(filepath: str = TEST_DATA_EXCEL) -> List[Dict]:
    return _sheet_to_dicts(filepath, "TestCases")


def update_test_case_status(tc_id: str, status: str,
                             filepath: str = TEST_DATA_EXCEL) -> None:
    """Write PASS/FAIL back into the Status column of the TestCases sheet."""
    try:
        import openpyxl
        if not os.path.exists(filepath):
            return
        wb = openpyxl.load_workbook(filepath)
        ws = wb["TestCases"]
        headers = [cell.value for cell in ws[1]]
        tc_col     = headers.index("TC ID")    + 1
        status_col = headers.index("Status")   + 1
        for row in ws.iter_rows(min_row=2):
            if row[tc_col - 1].value == tc_id:
                row[status_col - 1].value = status
                break
        wb.save(filepath)
    except Exception:
        pass          # non-fatal; test run continues
